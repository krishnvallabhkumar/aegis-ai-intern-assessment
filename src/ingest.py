"""Multi-format Aegis ingestion pipeline.
Supports PDF, scanned PDF, HTML, XLSX, DOCX, PNG/JPG, JSON and PPTX.
Deterministic extraction is preferred; OCR is used when text extraction is absent/too short.
Graphical topology is represented through a reviewed diagram-facts sidecar because OCR cannot
reliably recover line semantics.
"""
from pathlib import Path
import json, re, sys, hashlib
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from PIL import Image
import pytesseract
from bs4 import BeautifulSoup

SUPPORTED = {".pdf",".html",".htm",".xlsx",".docx",".png",".jpg",".jpeg",".json",".pptx"}

def sha256(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for b in iter(lambda:f.read(1024*1024),b""): h.update(b)
    return h.hexdigest()

def pdf_extract(path):
    reader=PdfReader(path)
    pages=[]
    for i,p in enumerate(reader.pages,1):
        txt=(p.extract_text() or "").strip()
        method="pdf_text"
        if len(re.sub(r"\s+","",txt)) < 40:
            try:
                from pdf2image import convert_from_path
                imgs=convert_from_path(str(path),dpi=180,first_page=i,last_page=i)
                txt=pytesseract.image_to_string(imgs[0])
                method="ocr"
            except Exception:
                method="pdf_text_failed_ocr"
        pages.append({"location":f"p{i}","text":txt,"method":method})
    return pages

def extract(path):
    ext=path.suffix.lower()
    if ext==".pdf": return pdf_extract(path)
    if ext in {".png",".jpg",".jpeg"}:
        return [{"location":"image","text":pytesseract.image_to_string(Image.open(path)),"method":"ocr"}]
    if ext in {".html",".htm"}:
        soup=BeautifulSoup(path.read_text(encoding="utf-8",errors="ignore"),"html.parser")
        return [{"location":"HTML","text":soup.get_text("\n"),"method":"html_text"}]
    if ext==".json":
        return [{"location":"JSON","text":json.dumps(json.loads(path.read_text()),indent=2),"method":"json_parse"}]
    if ext==".docx":
        d=Document(path); blocks=[]
        for i,p in enumerate(d.paragraphs,1):
            if p.text.strip(): blocks.append(f"[paragraph {i}] {p.text}")
        for ti,t in enumerate(d.tables,1):
            for ri,row in enumerate(t.rows,1):
                blocks.append(f"[table {ti} row {ri}] " + " | ".join(c.text for c in row.cells))
        return [{"location":"DOCX","text":"\n".join(blocks),"method":"docx_parse"}]
    if ext==".xlsx":
        wb=load_workbook(path,data_only=True); blocks=[]
        for ws in wb.worksheets:
            blocks.append(f"[SHEET {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                blocks.append(" | ".join("" if c is None else str(c) for c in row))
        return [{"location":"XLSX","text":"\n".join(blocks),"method":"xlsx_parse"}]
    if ext==".pptx":
        prs=Presentation(path); blocks=[]
        for i,s in enumerate(prs.slides,1):
            blocks.append(f"[SLIDE {i}]")
            for sh in s.shapes:
                if hasattr(sh,"text") and sh.text.strip(): blocks.append(sh.text)
        return [{"location":"PPTX","text":"\n".join(blocks),"method":"pptx_parse"}]
    return []

def main():
    if len(sys.argv)<2:
        print("Usage: python src/ingest.py <aegis-dataset>")
        raise SystemExit(2)
    root=Path(sys.argv[1])
    out=Path("outputs/raw_documents.jsonl"); out.parent.mkdir(exist_ok=True)
    rows=[]
    for p in sorted(root.rglob("*")):
        if p.is_file() and p.suffix.lower() in SUPPORTED:
            parts=extract(p)
            rows.append({"source":str(p.relative_to(root)),"sha256":sha256(p),"parts":parts})
    with out.open("w",encoding="utf-8") as f:
        for r in rows: f.write(json.dumps(r,ensure_ascii=False)+"\n")
    print(f"Ingested {len(rows)} files -> {out}")

if __name__=="__main__": main()
